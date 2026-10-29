"""Offline checks for preparing and converting manual-review workbooks.

Use small temporary Excel files to check metadata copying, dataset-ID
fields, manual corrections and preservation of existing annotations.
These tests do not collect publications or assign reuse categories.
"""
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from build_review_workbook import (
    attach_metadata_sheet_to_review_workbook,
    build_review_workbook,
    migrate_existing_review_workbook,
    numeric_study_ids,
    study_id_sort_key,
    load_mtbls_metadata,
)


class StudyIdDerivationTests(unittest.TestCase):
    def test_transposed_collector_metadata_becomes_review_snapshot(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source = temp_path / "search.xlsx"
            metadata_path = temp_path / "metadata.xlsx"
            output = temp_path / "review.xlsx"
            pd.DataFrame([{
                "Match.URL": "https://example.test/MTBLS42",
                "Match.Title": "Study of MTBLS42",
                "Match.Abstract": "MTBLS42",
            }]).to_excel(source, sheet_name="Combined", index=False)
            pd.DataFrame({"MTBLS42": {
                "MTBLS.Accession": "MTBLS42",
                "MTBLS.ReleaseDate": "2025-01-01",
                "MTBLS.Platforms": "LC-MS",
            }}).to_excel(metadata_path)

            frame, source_sheet = load_mtbls_metadata(metadata_path)
            self.assertEqual(source_sheet, "Sheet1 (transposed)")
            self.assertEqual(frame.loc[0, "MTBLS.Platforms"], "LC-MS")
            build_review_workbook(source, "Combined", output, metadata_xlsx=metadata_path)
            book = load_workbook(output, data_only=True)
            self.assertEqual(book["MTBLS_metadata"]["B2"].value, "MTBLS42")
            provenance = dict(book["Provenance"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(provenance["MTBLS metadata source worksheet"], "Sheet1 (transposed)")

    def test_review_automatically_embeds_collector_metadata_sheet(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source = temp_path / "search.xlsx"
            output = temp_path / "review.xlsx"
            with pd.ExcelWriter(source) as writer:
                pd.DataFrame([{
                    "Match.URL": "https://example.test/MTBLS42",
                    "Match.Title": "Study of MTBLS42",
                    "Match.Abstract": "MTBLS42",
                }]).to_excel(writer, sheet_name="Combined", index=False)
                pd.DataFrame([{
                    "MTBLS.Accession": "MTBLS42",
                    "MTBLS.ReleaseDate": "2025-01-01",
                }]).to_excel(writer, sheet_name="MTBLS_metadata", index=False)
            build_review_workbook(source, "Combined", output)
            book = load_workbook(output, data_only=True)
            self.assertEqual(book["MTBLS_metadata"]["A2"].value, "MTBLS42")
            provenance = dict(book["Provenance"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(provenance["MTBLS metadata source workbook"], "search.xlsx")


    def test_numeric_ids_preserve_repository_number_collisions(self):
        study_ids = "MTBLS1, MSV000095095, ST000001"

        self.assertEqual(numeric_study_ids(study_ids), "1; 95095; 1")
        self.assertEqual(study_id_sort_key(study_ids), 1)

    def test_review_workbook_has_reproducible_study_id_contract(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source = temp_path / "source.xlsx"
            output = temp_path / "review.xlsx"
            pd.DataFrame([
                {
                    "Match.URL": "https://example.test/MTBLS42",
                    "Match.Title": "A study referencing MTBLS42 and ST000007",
                    "Match.Abstract": "MTBLS42; ST000007",
                    "Match.Journal": "Journal",
                    "Match.Year": 2026,
                    "Match.Authors": "Author",
                }
            ]).to_excel(source, sheet_name="Combined", index=False)

            build_review_workbook(source, "Combined", output)

            workbook = load_workbook(output, data_only=False)
            review = workbook["Review"]
            headers = [cell.value for cell in review[1]]
            self.assertIn("StudyIDs (detected)", headers)
            self.assertIn("StudyIDs (corrected)", headers)
            self.assertIn("StudyIDs (numeric)", headers)
            self.assertIn("StudyID sort key", headers)
            self.assertIn("Reviewer.Decision.Timestamp (UTC)", headers)
            values = dict(zip(headers, [cell.value for cell in review[2]]))
            self.assertEqual(values["StudyIDs (detected)"], "MTBLS42, ST000007")
            self.assertEqual(values["StudyIDs (corrected)"], None)
            self.assertEqual(values["StudyIDs (numeric)"], "42; 7")
            self.assertEqual(values["StudyID sort key"], 7)
            self.assertIn("Provenance", workbook.sheetnames)
            provenance = {
                field: value
                for field, value in workbook["Provenance"].iter_rows(min_row=2, values_only=True)
            }
            self.assertIn("Input SHA-256", provenance)
            self.assertIn("Generator SHA-256", provenance)
            self.assertEqual(provenance["StudyID source columns"], "Match.Title; Match.Abstract; Match.URL")
            self.assertFalse(any(
                isinstance(cell.value, str) and "NOW(" in cell.value
                for row in review.iter_rows()
                for cell in row
            ))

    def test_legacy_migration_preserves_annotations_and_makes_corrections_sparse(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source = temp_path / "legacy.xlsx"
            output = temp_path / "reformatted.xlsx"
            pd.DataFrame([
                {
                    "URL": "https://example.test/1",
                    "Title": "First title",
                    "Abstract (truncated)": "Abstract",
                    "Journal": "Journal",
                    "Year": 2025,
                    "Authors": "Author",
                    "StudyIDs": "MTBLS42, ST000007",
                    "StudyIDs (corrected)": "MTBLS42, ST000007",
                    "Reviewer.Primary (pick)": "Method or Tool Development",
                    "Reviewer.Subcats (free text)": "Workflow",
                    "Reviewer.Domain (pick)": "Metabolomics",
                    "Reviewer.Initials": "AB",
                    "Reviewer.Notes (e.g., how the data was used)": "Reuse",
                },
                {
                    "URL": "https://example.test/2",
                    "Title": "Second title",
                    "Abstract (truncated)": "Abstract",
                    "Journal": "Journal",
                    "Year": 2025,
                    "Authors": "Author",
                    "StudyIDs": "MTBLS88",
                    "StudyIDs (corrected)": " MTBLS881 ",
                    "Reviewer.Primary (pick)": None,
                    "Reviewer.Subcats (free text)": None,
                    "Reviewer.Domain (pick)": None,
                    "Reviewer.Initials": None,
                    "Reviewer.Notes (e.g., how the data was used)": "False Positive",
                },
            ]).to_excel(source, sheet_name="Review", index=False)

            migrate_existing_review_workbook(source, "Review", output)

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Instructions", "Provenance", "Review", "Review.CollapsedByStudy"])
            review = workbook["Review"]
            headers = [cell.value for cell in review[1]]
            first = dict(zip(headers, [cell.value for cell in review[2]]))
            second = dict(zip(headers, [cell.value for cell in review[3]]))
            self.assertIsNone(first["StudyIDs (corrected)"])
            self.assertEqual(second["StudyIDs (corrected)"], "MTBLS881")
            self.assertEqual(first["StudyIDs (numeric)"], "42; 7")
            self.assertEqual(first["StudyID sort key"], 7)
            self.assertEqual(first["Primary (pick)"], "Method or Tool Development")
            self.assertEqual(first["Reviewer.Notes"], "Reuse")
            self.assertEqual(second["Reviewer.Notes"], "False Positive")
            provenance = dict(workbook["Provenance"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(provenance["Correction migration"], "Retained only values that differed from the corresponding detected StudyIDs; surrounding whitespace was removed.")

    def test_metadata_snapshot_can_be_attached_without_rewriting_review(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source = temp_path / "source.xlsx"
            review = temp_path / "review.xlsx"
            metadata = temp_path / "metadata.xlsx"
            pd.DataFrame([
                {
                    "Match.URL": "https://example.test/MTBLS42",
                    "Match.Title": "A study",
                    "Match.Abstract": "MTBLS42",
                    "Match.Journal": "Journal",
                    "Match.Year": 2026,
                }
            ]).to_excel(source, sheet_name="Combined", index=False)
            pd.DataFrame([
                {"MTBLS.Accession": "MTBLS42", "MTBLS.ReleaseDate": "2025-01-01"}
            ]).to_excel(metadata, sheet_name="MTBLS_metadata", index=False)

            build_review_workbook(source, "Combined", review)
            attach_metadata_sheet_to_review_workbook(review, metadata)

            workbook = load_workbook(review, data_only=False)
            self.assertIn("MTBLS_metadata", workbook.sheetnames)
            self.assertEqual(workbook["MTBLS_metadata"]["A2"].value, "MTBLS42")
            self.assertEqual(workbook["Review"]["B2"].value, "A study")
            provenance = dict(workbook["Provenance"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(provenance["MTBLS metadata source workbook"], "metadata.xlsx")
            self.assertEqual(provenance["MTBLS metadata snapshot dimensions"], "1 data rows × 2 columns")


if __name__ == "__main__":
    unittest.main()
