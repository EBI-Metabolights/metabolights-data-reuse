#!/usr/bin/env python3
"""Prepare an Excel workbook for manual publication screening and classification.

Read a search-results sheet, copy publication details and find dataset
identifiers. Leave decisions, categories and reviewer fields blank for a
person to complete. Include source-file checksums, a grouped browsing sheet
and dataset metadata when available. No reuse labels are predicted.
"""
from __future__ import annotations
import argparse
from copy import copy
import hashlib
import pandas as pd
from pathlib import Path
from typing import List, Optional
from xlsxwriter.utility import xl_rowcol_to_cell

import re


STUDY_ID_PATTERN = re.compile(r"\b(?P<prefix>MTBLS|ST|MSV)(?P<number>\d+)\b", re.IGNORECASE)

def extract_study_ids_row(row, candidate_cols):
    """Find distinct repository accessions in the selected cells and return sorted text."""
    patt_mtbls = re.compile(r"\bMTBLS\d+\b", re.IGNORECASE)
    patt_mw = re.compile(r"\bST\d{6}\b", re.IGNORECASE)
    patt_massive = re.compile(r"\bMSV\d{7,10}\b", re.IGNORECASE)
    accs = set()
    for col in candidate_cols:
        if col is None or col not in row.index:
            continue
        val = str(row.get(col, '') or '')
        if not val.strip():
            continue
        for patt in (patt_mtbls, patt_mw, patt_massive):
            for m in patt.findall(val):
                accs.add(m.upper())
    return ", ".join(sorted(accs))


def numeric_study_ids(study_ids: object) -> str:
    """Return each detected accession's numeric part in the order supplied.

    Keep repeated numbers from different repositories: MTBLS1 and ST000001
    both contribute 1. These numbers are sorting aids; the full accession
    must be retained to identify its repository and dataset.
    """
    return "; ".join(str(int(match.group("number"))) for match in STUDY_ID_PATTERN.finditer(str(study_ids or "")))


def study_id_sort_key(study_ids: object) -> Optional[int]:
    """Return the smallest accession number for sorting, or None when no ID is found."""
    numbers = [int(match.group("number")) for match in STUDY_ID_PATTERN.finditer(str(study_ids or ""))]
    return min(numbers) if numbers else None


def sha256_file(path: Path) -> str:
    """Calculate a SHA-256 checksum so a later reader can identify the same file."""
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


PRIMARY_OPTIONS = [
    "Secondary Biological Analysis",
    "Method or Tool Development",
    "Data Integration or Meta-analysis",
    "Educational Use or Protocol Illustration",
]

def detect_col(df: pd.DataFrame, preferred: List[str], contains_any: Optional[List[str]] = None) -> Optional[str]:
    """Find a preferred column name without regard to letter case.

    If no preferred name exists, try the supplied words within column names.
    Return None when neither search finds a column.
    """
    cols_lc = {c.lower(): c for c in df.columns}
    for cand in preferred:
        if cand.lower() in cols_lc: return cols_lc[cand.lower()]
    if contains_any:
        for c in df.columns:
            if any(tok in c.lower() for tok in contains_any): return c
    return None


def load_mtbls_metadata(path: Path) -> tuple[pd.DataFrame, str]:
    """Read metadata with one accession per row and check that accessions are unique.

    Accept an MTBLS_metadata sheet, or the original Sheet1 export with one
    accession per column, which is transposed before use. Return the table
    and the source-sheet description.
    """
    with pd.ExcelFile(path) as workbook:
        if "MTBLS_metadata" in workbook.sheet_names:
            metadata = pd.read_excel(workbook, sheet_name="MTBLS_metadata")
            source_sheet = "MTBLS_metadata"
        elif "Sheet1" in workbook.sheet_names:
            raw = pd.read_excel(workbook, sheet_name="Sheet1", index_col=0)
            if "MTBLS.Accession" not in raw.index:
                raise ValueError("Sheet1 is not a transposed MTBLS metadata export")
            metadata = raw.T.reset_index().rename(columns={"index": "Unnamed: 0"})
            source_sheet = "Sheet1 (transposed)"
        else:
            raise ValueError("Metadata workbook has neither MTBLS_metadata nor transposed Sheet1")
    if "MTBLS.Accession" not in metadata or metadata["MTBLS.Accession"].isna().any():
        raise ValueError("Metadata must contain a nonempty MTBLS.Accession for every row")
    if metadata["MTBLS.Accession"].duplicated().any():
        raise ValueError("Metadata contains duplicate MTBLS.Accession values")
    return metadata, source_sheet

def build_review_workbook(in_path: Path, in_sheet: str, out_path: Path,
                          title_col_override: Optional[str] = None, abstract_col_override: Optional[str] = None,
                          url_col_override: Optional[str] = None, journal_col_override: Optional[str] = None,
                          year_col_override: Optional[str] = None,
                          detected_study_ids_override: Optional[pd.Series] = None,
                          metadata_xlsx: Optional[Path] = None) -> None:
    """Create the manual-review workbook from the selected input sheet.

    Copy publication details, derive detected accessions and leave corrected
    IDs and classification fields blank. Embed metadata from metadata_xlsx or,
    when available, the input workbook's MTBLS_metadata sheet. The grouped sheet
    is a browsing aid, not the source for publication or dataset counts.
    """
    df = pd.read_excel(in_path, sheet_name=in_sheet)
    metadata = None
    metadata_source_sheet = None
    metadata_source = metadata_xlsx
    if metadata_source is None:
        with pd.ExcelFile(in_path) as input_book:
            if "MTBLS_metadata" in input_book.sheet_names:
                metadata_source = in_path
    if metadata_source is not None:
        metadata, metadata_source_sheet = load_mtbls_metadata(metadata_source)

    title_col = title_col_override or detect_col(df, ["Match.Title","Title"], contains_any=["title"])
    abstract_col = abstract_col_override or detect_col(df, ["Match.Abstract","Abstract","Summary"], contains_any=["abstract","summary"])
    url_col = url_col_override or detect_col(df, ["URL","Match.URL","Link"], contains_any=["url","link"])
    journal_col = journal_col_override or detect_col(df, ["Match.Journal","Journal","Source Title","Publication Name"], contains_any=["journal","source"])
    year_col = year_col_override or detect_col(df, ["Match.Year","Year","Publication Year","PubYear"], contains_any=["year"])
    doi_col = detect_col(df, ["Match.DOI", "DOI"])
    pmcid_col = detect_col(df, ["Match.PMCID", "PMCID"])
    authors_col = detect_col(df, ["Authors","Author(s)","Author Names","Author list"], contains_any=["author"])

    review = pd.DataFrame()
    if url_col: review["URL"] = df[url_col]
    if doi_col: review["DOI"] = df[doi_col]
    if pmcid_col: review["PMCID"] = df[pmcid_col]
    if title_col: review["Title"] = df[title_col]
    if abstract_col:
        review["Abstract (truncated)"] = df[abstract_col].astype(str).str.replace(r"\s+"," ", regex=True).str.slice(0,500)

    if journal_col: review["Journal"] = df[journal_col]
    if year_col: review["Year"] = df[year_col]
    if authors_col: review["Authors"] = df[authors_col]

    # Search publication text, links and columns likely to contain dataset accessions.
    candidate_cols = []
    if title_col: candidate_cols.append(title_col)
    if abstract_col: candidate_cols.append(abstract_col)
    if url_col: candidate_cols.append(url_col)
    for c in df.columns:
        cl = c.lower()
        if any(tok in cl for tok in ["study","accession","id","massive","gnps","metabolights","workbench","repository","record","url"]):
            candidate_cols.append(c)
    candidate_cols = list(dict.fromkeys(candidate_cols))
    if detected_study_ids_override is None:
        review["StudyIDs (detected)"] = df.apply(lambda r: extract_study_ids_row(r, candidate_cols), axis=1)
    else:
        if len(detected_study_ids_override) != len(df):
            raise ValueError("detected_study_ids_override must have one value per input row")
        review["StudyIDs (detected)"] = detected_study_ids_override.fillna("").astype(str).to_list()
    # Keep detected IDs unchanged; the reviewer records corrections in a separate field.
    review["StudyIDs (corrected)"] = ""
    review["StudyIDs (numeric)"] = review["StudyIDs (detected)"].map(numeric_study_ids)
    review["StudyID sort key"] = review["StudyIDs (detected)"].map(study_id_sort_key)

    # Leave screening decisions and classification fields for the reviewer to complete.
    review["Primary (pick)"] = ""
    review["Subcategories (free text)"] = ""
    review["Domain (pick)"] = ""
    review["Reviewer.Initials"] = ""
    review["Reviewer.Decision.Timestamp (UTC)"] = ""
    review["Reviewer.Notes"] = ""
    review["Status"] = ""

    with pd.ExcelWriter(out_path, engine="xlsxwriter") as writer:
        wb = writer.book
        wrap_fmt = wb.add_format({"text_wrap": True, "valign": "top"})
        # Explain the manual fields and the limits of the progress indicator.
        instr = pd.DataFrame({"How to use":[
            "Use filters on 'Review'. Assign Primary, Subcategories, Domain, and reviewer initials manually. Record decision timestamps manually in ISO 8601 UTC format (for example, 2026-10-16T14:30:00Z).",
            "Status is Reviewed when Primary (pick) is nonempty, otherwise Pending. This is a progress marker, not a check that screening is complete.",
            "Keep the automatically detected dataset IDs unchanged. Record manual edits only in StudyIDs (corrected).",
            "StudyIDs (numeric) lists the numeric part of every detected accession. StudyID sort key is the smallest such number; it is intended only for sorting and does not identify a repository.",
            "DOI and PMCID are imported from the search record for audit and analysis. A different DOI is not automatically grouped with a preprint or another version; reviewer decisions remain manual.",
            "Review.CollapsedByStudy groups records that mention the same set of detected study IDs.",
            "Keep MTBLS_metadata as the saved repository metadata used for modality and release-year analyses. The sheet is not technically locked."
        ]})
        instr.to_excel(writer, sheet_name="Instructions", index=False)
        ws_i = writer.sheets["Instructions"]; ws_i.set_column(0,0,110)

        provenance = pd.DataFrame([
            ("Generator", Path(__file__).name),
            ("Generator SHA-256", sha256_file(Path(__file__))),
            ("Input workbook", in_path.name),
            ("Input worksheet", in_sheet),
            ("Input SHA-256", sha256_file(in_path)),
            ("Input rows", len(df)),
            ("StudyID source columns", "; ".join(candidate_cols)),
            ("StudyID extraction", "Distinct MTBLS, ST, and MSV accessions found in the selected columns; uppercase and sorted as text."),
            ("Numeric IDs", "Numbers from each detected accession, in the same order and separated by semicolons."),
            ("Sort key", "Smallest numeric accession component; use with the detected accession as a tie-breaker."),
            ("Correction policy", "StudyIDs (corrected) starts blank and is reserved for documented manual review; detected values are never replaced."),
            ("Decision timestamp policy", "Enter the decision time manually as UTC text, for example 2026-10-16T14:30:00Z; recalculation does not change it."),
            *([] if metadata is None else [
                ("MTBLS metadata source workbook", metadata_source.name),
                ("MTBLS metadata source worksheet", metadata_source_sheet),
                ("MTBLS metadata source SHA-256", sha256_file(metadata_source)),
                ("MTBLS metadata rows", len(metadata)),
            ]),
        ], columns=["Field", "Value"])
        provenance.to_excel(writer, sheet_name="Provenance", index=False)
        ws_p = writer.sheets["Provenance"]
        ws_p.set_column(0, 0, 24)
        ws_p.set_column(1, 1, 110, wrap_fmt)

        # Write one review row for each input search-result row.
        review.to_excel(writer, sheet_name="Review", index=False)
        ws = writer.sheets["Review"]
        header_fmt = wb.add_format({"bold": True, "text_wrap": True, "valign": "top", "bg_color": "#EEE"})
        done_fmt = wb.add_format({"bg_color": "#E2F0D9"})
        cols = {c:i for i,c in enumerate(review.columns)}
        for name,width,fmt in [("URL",18,wrap_fmt),("Title",40,wrap_fmt),("Abstract (truncated)",70,wrap_fmt)]:
            if name in cols: ws.set_column(cols[name], cols[name], width, fmt)
        for name in ["DOI", "PMCID", "Journal", "Year", "StudyIDs (detected)", "StudyIDs (corrected)", "StudyIDs (numeric)", "StudyID sort key", "Authors", "Status"]:
            if name in cols: ws.set_column(cols[name], cols[name], 18)
        for name in ["Primary (pick)", "Subcategories (free text)", "Domain (pick)", "Reviewer.Initials", "Reviewer.Decision.Timestamp (UTC)", "Reviewer.Notes"]:
            if name in cols: ws.set_column(cols[name], cols[name], 24, wrap_fmt)

        ws.set_row(0, None, header_fmt)
        freeze_col_idx = 2 if "URL" in cols else 1
        ws.freeze_panes(1, freeze_col_idx)
        ws.autofilter(0, 0, len(review), len(review.columns)-1)

        # Add category menus and a status formula based only on the primary-category field.
        primary_options = [*PRIMARY_OPTIONS, "None"]
        domain_options = ["Metabolomics","Lipidomics","Both","Other"]
        if "Primary (pick)" in cols:
            i = cols["Primary (pick)"]; ws.data_validation(1,i,len(review),i,{"validate":"list","source":primary_options})
        if "Domain (pick)" in cols:
            i = cols["Domain (pick)"]; ws.data_validation(1,i,len(review),i,{"validate":"list","source":domain_options})
        if "Reviewer.Initials" in cols:
            i = cols["Reviewer.Initials"]; ws.data_validation(1,i,len(review),i,{"validate":"text length","criteria":">=","value":2})
            ws.data_validation(1,i,len(review),i,{"validate":"text length","criteria":"<=","value":4})
        if "Status" in cols and "Primary (pick)" in cols:
            st_c = cols["Status"]; pick_c = cols["Primary (pick)"]
            for r in range(1, len(review)+1):
                ws.write_formula(r, st_c, f'=IF({xl_rowcol_to_cell(r,pick_c)}<>"","Reviewed","Pending")')
        if "Status" in cols:
            ws.conditional_format(1, cols["Status"], len(review), cols["Status"], {
                "type":"text","criteria":"containing","value":"Reviewed","format":done_fmt})
        # Group identical sets of detected accessions for browsing, not for analysis counts.
        def norm_ids(s):
            """Sort distinct accessions so rows listing the same set can be grouped for browsing."""
            toks = sorted(set([t.strip().upper() for t in str(s or "").split(",") if t.strip()]))
            return "; ".join(toks)
        tmp = review.copy()
        tmp["StudyIDs_norm"] = tmp.get("StudyIDs (detected)", "").apply(norm_ids)
        has_ids = tmp["StudyIDs_norm"].astype(str).str.len() > 0
        with_ids = tmp[has_ids].copy()
        no_ids = tmp[~has_ids].copy()

        def agg_union(series):
            """Combine distinct comma- or semicolon-separated labels in a browsing group."""
            vals = set()
            for v in series.dropna().astype(str):
                for tok in [t.strip() for t in re.split(r"[;,]", v) if t.strip()]:
                    vals.add(tok)
            return ", ".join(sorted(vals))

        agg_map = {}
        for col in review.columns:
            if col in ["Primary (pick)","Subcategories (free text)","Domain (pick)"]:
                agg_map[col] = agg_union
            elif col in ["URL", "DOI", "PMCID", "Title", "Abstract (truncated)", "Journal", "Year", "StudyIDs (detected)", "StudyIDs (corrected)", "StudyIDs (numeric)", "StudyID sort key", "Authors"]:
                agg_map[col] = "first"
            elif col.startswith("Reviewer.") or col == "Status":
                agg_map[col] = "first"
            else:
                agg_map[col] = "first"

        collapsed = with_ids.groupby("StudyIDs_norm", as_index=False).agg(agg_map)
        collapsed["StudyIDs (detected)"] = collapsed["StudyIDs_norm"]
        if "StudyIDs_norm" in collapsed.columns:
            collapsed = collapsed.drop(columns=["StudyIDs_norm"])
        sizes = with_ids.groupby("StudyIDs_norm").size().rename("Collapsed.FromN")
        collapsed = collapsed.merge(sizes, left_on="StudyIDs (detected)", right_index=True, how="left")
        collapsed.to_excel(writer, sheet_name="Review.CollapsedByStudy", index=False)
        ws_c = writer.sheets["Review.CollapsedByStudy"]
        ws_c.autofilter(0, 0, len(collapsed), len(collapsed.columns)-1)
        ws_c.freeze_panes(1, 1)

        if len(no_ids) > 0:
            no_ids.to_excel(writer, sheet_name="Review.NoStudyID", index=False)
            ws_n = writer.sheets["Review.NoStudyID"]
            ws_n.autofilter(0, 0, len(no_ids), len(no_ids.columns)-1)
            ws_n.freeze_panes(1, 1)

        if metadata is not None:
            metadata.to_excel(writer, sheet_name="MTBLS_metadata", index=False)
            ws_m = writer.sheets["MTBLS_metadata"]
            ws_m.freeze_panes(1, 1)
            ws_m.autofilter(0, 0, len(metadata), len(metadata.columns) - 1)
            ws_m.set_column(0, len(metadata.columns) - 1, 24, wrap_fmt)


def attach_metadata_sheet_to_review_workbook(
    review_xlsx: Path,
    metadata_xlsx: Path,
    metadata_sheet: str = "MTBLS_metadata",
) -> None:
    """Copy a metadata sheet, its formatting and its source checksum into a review.

    Require the metadata sheet to be absent from the destination and require
    an existing Provenance sheet for recording the source information.
    """
    from openpyxl import load_workbook

    review_book = load_workbook(review_xlsx)
    if metadata_sheet in review_book.sheetnames:
        raise ValueError(f"Review workbook already has a {metadata_sheet!r} sheet")

    metadata_book = load_workbook(metadata_xlsx, data_only=False)
    if metadata_sheet not in metadata_book.sheetnames:
        raise ValueError(f"Metadata source has no {metadata_sheet!r} sheet")
    source = metadata_book[metadata_sheet]
    target = review_book.create_sheet(metadata_sheet)

    for row in source.iter_rows():
        for source_cell in row:
            target_cell = target.cell(source_cell.row, source_cell.column, source_cell.value)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)

    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))
    target.freeze_panes = source.freeze_panes
    target.auto_filter.ref = source.auto_filter.ref
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for index, dimension in source.row_dimensions.items():
        target.row_dimensions[index].height = dimension.height
        target.row_dimensions[index].hidden = dimension.hidden

    if "Provenance" not in review_book.sheetnames:
        raise ValueError("Review workbook has no Provenance sheet")
    provenance = review_book["Provenance"]
    provenance.append(["MTBLS metadata source workbook", metadata_xlsx.name])
    provenance.append(["MTBLS metadata source worksheet", metadata_sheet])
    provenance.append(["MTBLS metadata source SHA-256", sha256_file(metadata_xlsx)])
    provenance.append([
        "MTBLS metadata snapshot dimensions",
        f"{source.max_row - 1} data rows × {source.max_column} columns",
    ])
    metadata_book.close()
    review_book.save(review_xlsx)


def _first_present_column(frame: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    """Return the first supplied column name that exists, or None."""
    for candidate in candidates:
        if candidate in frame.columns:
            return candidate
    return None


def _text_or_empty(value: object) -> str:
    """Convert a populated value to text and a missing value to empty text."""
    return "" if pd.isna(value) else str(value)


def migrate_existing_review_workbook(in_path: Path, in_sheet: str, out_path: Path) -> None:
    """Convert an older reviewed workbook while preserving manual decisions.

    Older sheets repeated detected accessions in the corrected-ID column even
    when no edit was made. Leave that field blank when it merely repeats the
    detected value; preserve actual corrections and the reviewer's annotations.
    Regenerate numeric sorting fields from the detected accessions.
    """
    legacy = pd.read_excel(in_path, sheet_name=in_sheet)
    detected_col = _first_present_column(legacy, ["StudyIDs (detected)", "StudyIDs"])
    if detected_col is None:
        raise ValueError("The source review sheet has no StudyIDs column to migrate")
    corrected_col = _first_present_column(legacy, ["StudyIDs (corrected)"])

    build_review_workbook(
        in_path=in_path,
        in_sheet=in_sheet,
        out_path=out_path,
        title_col_override="Title",
        abstract_col_override="Abstract (truncated)",
        url_col_override="URL",
        journal_col_override="Journal",
        year_col_override="Year",
        detected_study_ids_override=legacy[detected_col],
    )

    from openpyxl import load_workbook

    workbook = load_workbook(out_path)
    review_sheet = workbook["Review"]
    output_cols = {cell.value: cell.column for cell in review_sheet[1]}
    annotation_sources = {
        "Primary (pick)": ["Primary (pick)", "Reviewer.Primary (pick)"],
        "Subcategories (free text)": ["Subcategories (free text)", "Reviewer.Subcats (free text)"],
        "Domain (pick)": ["Domain (pick)", "Reviewer.Domain (pick)"],
        "Reviewer.Initials": ["Reviewer.Initials"],
        "Reviewer.Decision.Timestamp (UTC)": ["Reviewer.Decision.Timestamp (UTC)", "Decision.Timestamp"],
        "Reviewer.Notes": ["Reviewer.Notes", "Reviewer.Notes (e.g., how the data was used)"],
    }
    annotation_cols = {
        output: _first_present_column(legacy, candidates)
        for output, candidates in annotation_sources.items()
    }

    for output_row, legacy_row in enumerate(legacy.itertuples(index=False, name=None), start=2):
        source_values = dict(zip(legacy.columns, legacy_row))
        detected = _text_or_empty(source_values[detected_col])
        historical_corrected = _text_or_empty(source_values[corrected_col]) if corrected_col else ""
        manual_corrected = ""
        if historical_corrected.strip() and historical_corrected.strip() != detected.strip():
            manual_corrected = historical_corrected.strip()
        review_sheet.cell(output_row, output_cols["StudyIDs (corrected)"], manual_corrected or None)
        for output_name, source_name in annotation_cols.items():
            if source_name is not None:
                value = _text_or_empty(source_values[source_name])
                review_sheet.cell(output_row, output_cols[output_name], value or None)

    instructions = workbook["Instructions"]
    instructions.delete_rows(2, instructions.max_row)
    instructions.append(["This workbook was converted from an older review format. Existing decisions and actual dataset-ID corrections were retained."])
    instructions.append(["StudyIDs (corrected) is blank when the older value merely repeated StudyIDs (detected)."])
    instructions.append(["StudyIDs (numeric) and StudyID sort key are calculated from StudyIDs (detected)."])
    instructions.append(["Record decision timestamps manually in ISO 8601 UTC format."])

    provenance = workbook["Provenance"]
    provenance.delete_rows(2, provenance.max_row)
    provenance_rows = [
        ("Generator", Path(__file__).name),
        ("Generator SHA-256", sha256_file(Path(__file__))),
        ("Reformatted source workbook", in_path.name),
        ("Reformatted source worksheet", in_sheet),
        ("Reformatted source SHA-256", sha256_file(in_path)),
        ("Input rows", len(legacy)),
        ("Detected StudyID source", detected_col),
        ("Historical corrected source", corrected_col or "None"),
        ("Correction migration", "Retained only values that differed from the corresponding detected StudyIDs; surrounding whitespace was removed."),
        ("Numeric IDs", "Numbers from each detected accession, in the same order and separated by semicolons."),
        ("Sort key", "Smallest numeric accession component; use with the detected accession as a tie-breaker."),
    ]
    for field, value in provenance_rows:
        provenance.append([field, value])

    workbook.save(out_path)

def parse_args():
    """Read the input sheet, output path and optional publication-column names."""
    p = argparse.ArgumentParser(description="Create an Excel workbook for manual screening and classification of publication-dataset matches.")
    p.add_argument("--in", dest="in_path", required=True)
    p.add_argument("--in-sheet", dest="in_sheet", required=True)
    p.add_argument("--out", dest="out_path", required=True)
    p.add_argument("--title-col", dest="title_col", default="Match.Title")
    p.add_argument("--abstract-col", dest="abstract_col", default="Match.Abstract")
    p.add_argument("--url-col", dest="url_col", default="Match.URL")
    p.add_argument("--journal-col", dest="journal_col", default="Match.Journal")
    p.add_argument("--year-col", dest="year_col", default="Match.Year")
    p.add_argument("--metadata-xlsx", type=Path,
                   help="Original transposed metadata export or a workbook with an MTBLS_metadata sheet; defaults to the input workbook when it has that sheet.")
    return p.parse_args()

def main():
    """Build a blank manual-review workbook using the command-line file paths."""
    args = parse_args()
    build_review_workbook(in_path=Path(args.in_path), in_sheet=args.in_sheet, out_path=Path(args.out_path),
                          title_col_override=args.title_col, abstract_col_override=args.abstract_col,
                          url_col_override=args.url_col, journal_col_override=args.journal_col,
                          year_col_override=args.year_col, metadata_xlsx=args.metadata_xlsx)

if __name__ == "__main__":
    main()
